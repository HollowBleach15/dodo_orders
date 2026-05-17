import os
from celery import shared_task
from django.conf import settings
from openpyxl import Workbook
from datetime import date

from .services import RevenueReport, ProductReport


@shared_task(bind=True)
def generate_full_report_xlsx(self, date_from: str, date_to: str) -> str:
    """Генерирует полный Excel-отчёт и сохраняет в MEDIA_ROOT/reports/."""
    df = date.fromisoformat(date_from)
    dt = date.fromisoformat(date_to)

    wb = Workbook()

    # Лист 1: Выручка по дням
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Выручка"
    ws1.append(["Дата", "Заказов", "Выручка, ₽"])
    for r in RevenueReport.revenue_by_day(df, dt):
        ws1.append([r["day"].strftime("%d.%m.%Y"), r["orders"], float(r["revenue"] or 0)])

    # Лист 2: По филиалам
    ws2 = wb.create_sheet("Филиалы")
    ws2.append(["Филиал", "Заказов", "Выручка", "Средний чек"])
    for r in RevenueReport.revenue_by_branch(df, dt):
        ws2.append([r["branch__name"], r["orders"],
                    float(r["revenue"] or 0), float(r["avg_check"] or 0)])

    # Лист 3: ТОП товаров
    ws3 = wb.create_sheet("ТОП товаров")
    ws3.append(["Товар", "Категория", "Продано", "Выручка"])
    for r in ProductReport.top_products(df, dt, limit=100):
        ws3.append([r["product__name"], r["product__category"],
                    r["qty_sold"], float(r["revenue"] or 0)])

    out_dir = os.path.join(settings.MEDIA_ROOT, "reports")
    os.makedirs(out_dir, exist_ok=True)
    filename = f"full_report_{date_from}_{date_to}_{self.request.id[:8]}.xlsx"
    path = os.path.join(out_dir, filename)
    wb.save(path)
    return f"{settings.MEDIA_URL}reports/{filename}"
