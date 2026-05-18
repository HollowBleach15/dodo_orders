# reports/export.py
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def order_to_excel(order) -> bytes:
    """Генерирует Excel-накладную для заказа. Возвращает bytes."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None, "Workbook должен иметь активный лист"
    ws.title = f"Заказ #{order.pk}"

    header_font = Font(bold=True, size=12)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Заголовок
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Накладная к заказу #{order.pk}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Мета-данные
    meta = [
        ("Клиент:", order.client.full_name),
        ("Телефон:", order.client.phone),
        ("Филиал:", order.branch.name),
        ("Тип:", order.get_order_type_display()),
        ("Статус:", order.get_status_display()),
        ("Дата:", order.created_at.strftime("%d.%m.%Y %H:%M")),
    ]
    for row_idx, (label, value) in enumerate(meta, start=2):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=str(value))

    # Таблица позиций
    header_row = len(meta) + 3
    headers = ["№", "Товар", "Кол-во", "Цена за ед.", "Сумма"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(order.items.select_related("product").all(), 1):
        r = header_row + i
        for col, val in enumerate(
            [i, item.product.name, item.quantity,
             float(item.price), float(item.line_total)], 1
        ):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = border

    # Итог
    total_row = header_row + order.items.count() + 1
    if order.discount_percent > 0:
        ws.cell(row=total_row, column=4, value="Скидка:").font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=f"{order.discount_percent}%")
        total_row += 1
    ws.cell(row=total_row, column=4, value="ИТОГО:").font = Font(bold=True, size=12)
    ws.cell(row=total_row, column=5, value=float(order.total)).font = Font(bold=True, size=12)

    # Ширина колонок
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()